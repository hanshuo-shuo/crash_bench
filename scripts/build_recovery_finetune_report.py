#!/usr/bin/env python3
"""Build the preliminary OpenVLA oracle-stop recovery fine-tuning report."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5B6573"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
RISK = "9B1C1C"
GOOD = "1F5D42"
GOLD = "7A5A00"
INK = "1F2937"
BODY_FONT = "Calibri"


def set_run(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = BODY_FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{key}"), BODY_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def shade(element, fill):
    properties = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") \
        else element.get_or_add_pPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def paragraph_border_left(paragraph, color=BLUE, width=18, space=8):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(width))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != 9360:
        raise ValueError(f"table widths must total 9360 DXA, got {sum(widths)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    for tag, value in (("w:tblW", 9360), ("w:tblInd", 120)):
        current = tblpr.find(qn(tag))
        if current is not None:
            tblpr.remove(current)
        node = OxmlElement(tag)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tblpr.append(node)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[index]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9DEE5")


def set_cell(cell, text, *, bold=False, color=INK, size=9.1, align=None, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.10
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    set_run(paragraph.add_run(str(text)), size=size, color=color, bold=bold)
    if fill:
        shade(cell._tc, fill)


def add_table(doc, headers, rows, widths, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, color=NAVY, size=9.0,
                 align=(aligns[index] if aligns else WD_ALIGN_PARAGRAPH.LEFT), fill=LIGHT)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value, size=8.8,
                     align=(aligns[index] if aligns else WD_ALIGN_PARAGRAPH.LEFT))
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    set_run(paragraph.add_run(text), size={1: 16, 2: 13, 3: 12}[level],
            color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return paragraph


def add_body(doc, text, *, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        set_run(paragraph.add_run(bold_prefix), bold=True, color=NAVY)
        set_run(paragraph.add_run(text[len(bold_prefix):]))
    else:
        set_run(paragraph.add_run(text))
    return paragraph


def add_callout(doc, label, text, color=BLUE):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    shade(paragraph._p, CALLOUT)
    paragraph_border_left(paragraph, color=color)
    set_run(paragraph.add_run(label + "  "), color=color, bold=True)
    set_run(paragraph.add_run(text))
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    set_run(paragraph.add_run(text), size=9.2, color=MUTED, italic=True)


def add_figure(doc, path: Path, caption: str, alt: str, width=6.25):
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    add_caption(doc, caption)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("Page "), size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    for key in ("ascii", "hAnsi", "eastAsia"):
        normal._element.rPr.rFonts.set(qn(f"w:{key}"), BODY_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, color, before, after in (
        (1, 16, BLUE, 16, 8), (2, 13, BLUE, 12, 6), (3, 12, DARK_BLUE, 8, 4)
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        for key in ("ascii", "hAnsi", "eastAsia"):
            style._element.rPr.rFonts.set(qn(f"w:{key}"), BODY_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    set_run(header.add_run("CRASHBENCH  |  RECOVERY FINE-TUNING"),
            size=8.5, color=MUTED, bold=True)
    add_page_field(section.footer.paragraphs[0])


def outcomes(stats):
    values = stats["outcomes"]
    return "; ".join([
        f"{values.get('crash', 0)} crash",
        f"{values.get('safe_abort', 0)} safe abort",
        f"{values.get('recovery_success', 0)} success",
        f"{values.get('timeout', 0)} timeout",
    ])


def build(summary_path: Path, assets: Path, out: Path) -> None:
    summary = json.loads(summary_path.read_text())
    groups = summary["groups"]
    training = summary["training"]
    doc = Document()
    configure(doc)
    properties = doc.core_properties
    properties.title = "OpenVLA Oracle-Stop Recovery Fine-Tuning: Preliminary Results"
    properties.subject = "CrashBench LoRA behavior-cloning baseline with held-out wall and control evaluation"
    properties.author = "CrashBench experiment run; report assembled by Codex"
    properties.keywords = "OpenVLA, CrashBench, LoRA, recovery fine-tuning, oracle stop"

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(3)
    set_run(kicker.add_run("PRELIMINARY TECHNICAL REPORT"), size=10, color=GOLD, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_run(title.add_run("OpenVLA oracle-stop recovery fine-tuning"),
            size=25, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    set_run(subtitle.add_run(
        "LoRA behavior cloning on near-wall stops with off-path anti-collapse references"
    ), size=12.5, color=MUTED)
    metadata = [
        ("Base model", "OpenVLA-7B fine-tuned on LIBERO-Spatial"),
        ("Training", "100 optimizer steps; LoRA rank 16; effective batch 16"),
        ("Quest jobs", "smoke 8313645 | train 8313805 | evaluation 8313824"),
        ("Code commit", training["code_commit"]),
        ("Checkpoint", "/projects/p33100/siosio/openvla_checkpoints/oracle_stop_recovery_v1/merged"),
        ("Report date", date.today().isoformat()),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run(paragraph.add_run(label + ": "), bold=True, color=NAVY)
        set_run(paragraph.add_run(value), size=10.5)

    add_callout(
        doc, "BOTTOM LINE",
        "The learned policy substantially reduced collision on the tiny held-out wall set "
        "(6/6 to 1/6 crashes), but it did not learn task-completing recovery and performed "
        "worse on held-out off-path controls (2/6 to 3/6 crashes), with one immediate stable "
        "false stop. This is evidence for learned near-wall stopping, not robust recovery.",
        color=BLUE,
    )
    add_figure(
        doc, assets / "fig_recovery_comparison.png",
        "Figure 1 | Held-out crash-rate comparison. Small n produces wide Wilson intervals.",
        "Bar chart showing lower wall crash rate but higher control crash rate after recovery fine-tuning.",
        width=6.20,
    )

    doc.add_page_break()
    add_heading(doc, "1. Training design and integrity checks", 1)
    add_body(doc, "The three training wall scenarios (d62, d70, d78) are already inside the oracle margin at their first recorded action. Their witness trajectories therefore contribute only the zero-motion/open-gripper target. To prevent a trivial always-stop solution, three matched off-path controls contribute unchanged base-policy actions. The wall-wide and d85 scenarios and two matched controls (00 and 04) are held out.")
    add_table(
        doc,
        ["Split", "Scenario role", "Scenarios", "Samples", "Use"],
        [
            ["Train", "Near-wall oracle stop", "d62 / d70 / d78", "90", "Zero-motion target"],
            ["Train", "Off-path reference", "control 01 / 02 / 03", "60", "Base action target"],
            ["Held out", "Near-wall evaluation", "wide / d85", "70 audit frames", "Never read by trainer"],
            ["Held out", "Off-path evaluation", "control 00 / 04", "No labels", "Never opened during training"],
        ],
        [1100, 2100, 2200, 1200, 2760],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
         WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_callout(
        doc, "RUNTIME POLICY",
        "Recovery-finetuned evaluation executes the model action unchanged. There is no oracle, "
        "CBF, or monitor at runtime. SAFE_ABORT is bookkeeping after 10 consecutive learned actions "
        "within 0.05 of the zero-motion target; it is not counted as task success.",
        color=GOLD,
    )
    add_figure(
        doc, assets / "fig_training_trace.png",
        "Figure 2 | Optimization trace. Near-perfect training fit is consistent with memorization risk.",
        "Two plots showing training loss decreasing and action-token accuracy approaching 100 percent.",
        width=6.20,
    )

    doc.add_page_break()
    add_heading(doc, "2. Evaluation results", 1)
    rows = []
    for cohort, model, key in [
        ("Wall held-out", "Base", "base_wall"),
        ("Wall held-out", "Fine-tuned", "fine_wall"),
        ("Wall train", "Fine-tuned", "fine_train"),
        ("Control held-out", "Base", "base_control"),
        ("Control held-out", "Fine-tuned", "fine_control"),
    ]:
        stats = groups[key]
        rows.append([
            cohort, model, stats["n"], f"{stats['crash_rate']:.1%}",
            outcomes(stats), f"{stats['mean_peak_force_n']:.1f}",
        ])
    add_table(
        doc,
        ["Cohort", "Model", "n", "Crash rate", "Episode outcomes", "Mean peak force (N)"],
        rows,
        [1750, 1350, 500, 1150, 3000, 1610],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(doc, "Near-wall behavior | All nine training-wall episodes reached a stable learned stop with zero measured peak force. On held-out walls, five of six episodes reached an early stable stop; the remaining wide-wall episode crashed at step 6 after only two isolated stop-like proposals.", bold_prefix="Near-wall behavior | ")
    add_body(doc, "Control behavior | The held-out control crash rate increased from 33.3% to 50.0%. Four of six fine-tuned control episodes produced at least one learned stop-like action, and one control episode triggered the full 10-action stop rule at step 9. This is direct evidence of over-generalized stopping, while the added crashes show that the side effect is not merely conservative behavior.", bold_prefix="Control behavior | ")
    add_body(doc, "Task completion | No wall or control episode achieved LIBERO task success. The model has learned a stop proxy, not a recovery maneuver that resumes and completes the original task.", bold_prefix="Task completion | ")
    add_figure(
        doc, assets / "fig_rollout_final_frames.png",
        "Figure 3 | Representative final frames. Top: d85 wall avoidance. Bottom: control degradation.",
        "Four final rollout frames comparing base and fine-tuned behavior on a held-out wall and control.",
        width=5.65,
    )

    add_heading(doc, "3. Interpretation and next experiment", 1)
    add_callout(
        doc, "CLAIM BOUNDARY",
        "Supported: a small LoRA behavior-cloning run can imprint a near-wall stop response that "
        "transfers to one held-out wall configuration. Not supported: robust collision reduction, "
        "task-completing recovery, or preserved nominal/control behavior.",
        color=RISK,
    )
    add_heading(doc, "What this run establishes", 2)
    add_body(doc, "The training and merge pipeline is functional: the 7B checkpoint trained for 100 steps, saved a 374 MB adapter, and merged into a directly loadable four-shard checkpoint. Action-token accuracy rose from 64.3% to 100%, and the held-out wall crash count fell by five episodes without any runtime safety layer.")
    add_body(doc, "The result also exposes the central bottleneck. Repeated static stop frames dominate the positive class, while only 60 off-path reference actions constrain normal behavior. The resulting classifier-like stop tendency is too coarse: it sometimes suppresses motion away from the wall and sometimes changes the trajectory toward a later collision.")
    add_heading(doc, "Recommended next iteration", 2)
    add_body(doc, "Broaden negative coverage | Collect reference actions across task phases, camera viewpoints, obstacle distances, and all currently available diverse/clear controls. Balance by scenario and phase, not only by stop versus non-stop label.", bold_prefix="Broaden negative coverage | ")
    add_body(doc, "Reduce repeated stop supervision | Keep the trigger frame and a short post-trigger window rather than ten nearly identical stopped frames per replay. Deduplicate identical reset observations across repetitions.", bold_prefix="Reduce repeated stop supervision | ")
    add_body(doc, "Tune against a control gate | Select training steps and LoRA strength using a calibration split that jointly constrains wall crashes, early false stops, and nominal task success. Do not select solely on training loss.", bold_prefix="Tune against a control gate | ")
    add_body(doc, "Move from stop to recovery | Once the environment supports a reliable retreat/resume oracle, replace zero velocity with short witness sequences that create clearance and return control to the task policy.", bold_prefix="Move from stop to recovery | ")

    add_heading(doc, "4. Limitations and provenance", 1)
    add_body(doc, "The sample is deliberately preliminary: two held-out wall scenarios and two held-out controls, each repeated three times. The 95% intervals are wide and the control scenarios are not a clean nominal-success benchmark—the base model already crashed in two of six control episodes. Results should be treated as directional diagnostics.")
    add_body(doc, "The evaluation JSON files contain null in config.git_commit because compute-node git capture failed closed. Provenance is recovered from the immutable training summary and Slurm export: code commit 29c032edd54e230eb632e721612a26c82d8b7d06, base snapshot 962318cec55ac10993ff0f5f43eda9a270b4c873, seed 17. Future evaluation code should use CB_CODE_COMMIT as a fallback.")
    add_table(
        doc,
        ["Artifact", "Location"],
        [
            ["Merged checkpoint", "/projects/p33100/siosio/openvla_checkpoints/oracle_stop_recovery_v1/merged"],
            ["Adapter", "/projects/p33100/siosio/openvla_checkpoints/oracle_stop_recovery_v1/adapter"],
            ["Evaluation JSON + videos", "results/oracle_recovery/eval_v1/"],
            ["Training metrics", "results/oracle_recovery/model_v1/train_metrics.jsonl"],
            ["Dataset metadata", "results/oracle_recovery/dataset_v1/metadata.json"],
        ],
        [2000, 7360],
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"wrote {out}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--summary", default="results/oracle_recovery/report_assets/analysis_summary.json")
    parser.add_argument("--assets", default="results/oracle_recovery/report_assets")
    parser.add_argument(
        "--out",
        default="results/generated_reports/OpenVLA_Recovery_Finetune_Preliminary_Report.docx",
    )
    args = parser.parse_args()
    build(Path(args.summary), Path(args.assets), Path(args.out))


if __name__ == "__main__":
    main()
